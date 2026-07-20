# -*- coding: utf-8 -*-
"""Genera el Manual de Usuario de Nexus ISP en formato Word (.docx)."""

from docx import Document
from docx.shared import Pt, RGBColor, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

# Paleta de colores
AZUL = RGBColor(0x1E, 0x3A, 0x8A)
AZUL_CLARO = RGBColor(0x25, 0x63, 0xEB)
GRIS = RGBColor(0x47, 0x55, 0x69)
BLANCO = RGBColor(0xFF, 0xFF, 0xFF)

doc = Document()

# --- Estilo base del documento ---
normal = doc.styles["Normal"]
normal.font.name = "Calibri"
normal.font.size = Pt(11)
normal.font.color.rgb = RGBColor(0x1F, 0x29, 0x37)


def set_cell_background(cell, color_hex):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), color_hex)
    tc_pr.append(shd)


def add_heading(text, level=1):
    p = doc.add_heading(level=level)
    run = p.add_run(text)
    run.font.name = "Calibri"
    if level == 1:
        run.font.size = Pt(18)
        run.font.color.rgb = AZUL
    elif level == 2:
        run.font.size = Pt(14)
        run.font.color.rgb = AZUL_CLARO
    else:
        run.font.size = Pt(12)
        run.font.color.rgb = GRIS
    return p


def add_para(text, bold=False, italic=False, color=None, size=11):
    p = doc.add_paragraph()
    run = p.add_run(text)
    run.bold = bold
    run.italic = italic
    run.font.size = Pt(size)
    if color:
        run.font.color.rgb = color
    return p


def add_bullet(text):
    return doc.add_paragraph(text, style="List Bullet")


def add_numbered(text):
    return doc.add_paragraph(text, style="List Number")


def add_note(text):
    p = doc.add_paragraph()
    p.paragraph_format.left_indent = Inches(0.2)
    run = p.add_run("ℹ  " + text)
    run.italic = True
    run.font.size = Pt(10)
    run.font.color.rgb = GRIS
    return p


def add_table(headers, rows):
    table = doc.add_table(rows=1, cols=len(headers))
    table.style = "Light Grid Accent 1"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    hdr = table.rows[0].cells
    for i, h in enumerate(headers):
        hdr[i].text = ""
        run = hdr[i].paragraphs[0].add_run(h)
        run.bold = True
        run.font.size = Pt(10)
        run.font.color.rgb = BLANCO
        set_cell_background(hdr[i], "1E3A8A")
    for row in rows:
        cells = table.add_row().cells
        for i, val in enumerate(row):
            cells[i].text = ""
            run = cells[i].paragraphs[0].add_run(val)
            run.font.size = Pt(10)
    return table


# ============================================================
# PORTADA
# ============================================================
for _ in range(6):
    doc.add_paragraph()

t = doc.add_paragraph()
t.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = t.add_run("MANUAL DE USUARIO")
r.font.size = Pt(34)
r.bold = True
r.font.color.rgb = AZUL

s = doc.add_paragraph()
s.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = s.add_run("Plataforma Nexus ISP")
r.font.size = Pt(22)
r.font.color.rgb = AZUL_CLARO

d = doc.add_paragraph()
d.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = d.add_run("Sistema de gestión de clientes, pagos y estado de servicio\npara proveedores de Internet (planes Starlink Residencial)")
r.font.size = Pt(12)
r.italic = True
r.font.color.rgb = GRIS

for _ in range(10):
    doc.add_paragraph()

v = doc.add_paragraph()
v.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = v.add_run("Versión 1.0")
r.font.size = Pt(11)
r.font.color.rgb = GRIS

doc.add_page_break()

# ============================================================
# 1. ¿QUÉ ES NEXUS ISP?
# ============================================================
add_heading("1. ¿Qué es Nexus ISP?", 1)
add_para("Plataforma Nexus ISP es un sistema web para administrar la operación de un proveedor de "
         "Internet. Con él usted puede:")
add_bullet("Registrar y editar clientes (datos de contacto, plan y tarifa).")
add_bullet("Registrar pagos mensuales de cada cliente.")
add_bullet("Ver de un vistazo el estado del negocio: clientes activos, suspendidos, pagos pendientes y total cobrado en el mes.")
add_bullet("Suspender y reactivar el servicio automáticamente según el pago de cada cliente.")
add_para("El sistema está pensado para ser sencillo y rápido de usar en el día a día.")

# ============================================================
# 2. ANTES DE EMPEZAR
# ============================================================
add_heading("2. Antes de empezar", 1)
add_para("Para usar la plataforma necesita:")
add_bullet("Un navegador web moderno (Chrome, Edge o Firefox actualizados).")
add_bullet("La dirección de acceso del sistema (por ejemplo, http://localhost:3000 si se ejecuta localmente, o la URL que le haya entregado el responsable técnico).")
add_bullet("Sus credenciales de acceso (usuario y contraseña).")
add_note("Nota de seguridad: las credenciales iniciales de fábrica son usuario 'admin' y contraseña "
         "'admin123'. Por seguridad, solicite al responsable técnico que las cambie antes de poner el "
         "sistema en producción.")

# ============================================================
# 3. INICIAR SESIÓN
# ============================================================
add_heading("3. Iniciar sesión", 1)
add_numbered("Abra el navegador y entre a la dirección del sistema.")
add_numbered("Verá la pantalla de inicio de sesión con el logo de Nexus.")
add_numbered("Escriba su usuario y su contraseña.")
add_numbered("Pulse el botón Iniciar sesión.")
add_para("¿Qué puede pasar?", bold=True)
add_bullet("Si los datos son correctos, entrará directamente al Dashboard.")
add_bullet("Si los datos son incorrectos, aparecerá un mensaje de error en rojo. Revise el usuario y la contraseña y vuelva a intentarlo.")
add_note("La sesión permanece activa durante 12 horas. Pasado ese tiempo deberá iniciar sesión nuevamente.")

# ============================================================
# 4. EL PANEL Y LA NAVEGACIÓN
# ============================================================
add_heading("4. El panel y la navegación", 1)
add_para("Una vez dentro, verá un menú con las tres secciones principales del sistema:")
add_table(
    ["Sección", "Para qué sirve"],
    [
        ["Dashboard", "Resumen general del negocio y pagos recientes."],
        ["Clientes", "Registrar, buscar y editar clientes."],
        ["Pagos", "Registrar pagos y ver el historial."],
    ],
)
add_para("En la parte superior también encontrará el botón para cerrar sesión.")

# ============================================================
# 5. DASHBOARD
# ============================================================
add_heading("5. Dashboard (Inicio)", 1)
add_para("El Dashboard es la pantalla de resumen. Muestra el estado actual del negocio en tiempo real.")
add_heading("Indicadores principales (tarjetas)", 2)
add_bullet("Clientes activos — total de clientes con servicio funcionando.")
add_bullet("Clientes suspendidos — total de clientes con el servicio cortado por falta de pago.")
add_bullet("Pagos pendientes — clientes que tienen un pago vencido o sin registrar.")
add_bullet("Cobrado este mes — suma total de los pagos recibidos durante el mes actual.")
add_heading("Clientes con pago pendiente", 2)
add_para("Una lista con los clientes que tienen pagos vencidos, para que pueda darles seguimiento.")
add_heading("Pagos recientes", 2)
add_para("Una tabla con los últimos pagos registrados (fecha, cliente, período y monto).")
add_note("Botón Actualizar: púlselo para recargar las cifras y ver la información más reciente.")

# ============================================================
# 6. GESTIÓN DE CLIENTES
# ============================================================
add_heading("6. Gestión de clientes", 1)
add_para("Entre a la sección Clientes desde el menú. La pantalla tiene tres áreas: el formulario para "
         "crear clientes, el buscador y la tabla de clientes.")

add_heading("6.1. Registrar un cliente nuevo", 2)
add_para("En el formulario de la izquierda complete los datos:")
add_table(
    ["Campo", "¿Obligatorio?", "Detalle"],
    [
        ["Nombre completo", "Sí", "Nombre del cliente."],
        ["Teléfono", "Sí", "Número de contacto."],
        ["Correo electrónico", "No", "Opcional."],
        ["Dirección", "Sí", "Ubicación del servicio."],
        ["Nombre del plan", "Sí", "Precargado como 'Plan Starlink Residencial'."],
        ["Tarifa mensual", "Sí", "Precargada en $40. Puede modificarla."],
        ["Primer mes gratis", "No", "Casilla opcional (ver abajo)."],
    ],
)
add_para("Pulse Guardar para crear el cliente.")
add_para("Sobre la casilla 'Primer mes gratis':", bold=True)
add_bullet("Activada: el cliente queda Activo y con un mes de servicio por delante.")
add_bullet("Desactivada: el cliente queda Suspendido desde el inicio, hasta que registre su primer pago.")

add_heading("6.2. Buscar y filtrar clientes", 2)
add_bullet("Buscar escribiendo el nombre, teléfono o ID del cliente.")
add_bullet("Filtrar por estado con el menú desplegable: Todos, Activos o Suspendidos.")
add_para("La cantidad de resultados se muestra para que sepa cuántos clientes coinciden.")

add_heading("6.3. La tabla de clientes", 2)
add_para("Muestra todos los clientes con: ID, Nombre, Teléfono, Plan, Estado del servicio "
         "(Activo / Suspendido) y un botón para Editar.")

add_heading("6.4. Editar un cliente", 2)
add_numbered("Pulse el botón Editar en la fila del cliente.")
add_numbered("Se abrirá una ventana con sus datos.")
add_numbered("Modifique lo que necesite: nombre, teléfono, correo, dirección, plan, tarifa o el estado del servicio (Activo / Suspendido).")
add_numbered("Pulse Guardar para confirmar, o Cancelar para descartar los cambios.")
add_note("Cambiar el estado a Activo o Suspendido manualmente le permite forzar el corte o la "
         "reactivación del servicio sin esperar a un pago.")

# ============================================================
# 7. GESTIÓN DE PAGOS
# ============================================================
add_heading("7. Gestión de pagos", 1)
add_para("Entre a la sección Pagos desde el menú.")

add_heading("7.1. Registrar un pago", 2)
add_para("En el formulario complete:")
add_table(
    ["Campo", "Detalle"],
    [
        ["Cliente", "Seleccione el cliente del menú desplegable."],
        ["Monto", "Cantidad pagada (precargado en $40)."],
        ["Fecha de pago", "Fecha en que se recibió el pago (por defecto, hoy)."],
        ["Período", "Mes que cubre el pago, en formato AAAA-MM (por defecto, el mes actual)."],
        ["Nota", "Comentario opcional (por ejemplo, 'pago en efectivo')."],
    ],
)
add_para("Pulse Guardar para registrar el pago.")
add_para("¿Qué ocurre al registrar un pago?", bold=True)
add_bullet("La fecha de vencimiento del cliente se extiende 30 días a partir de la fecha de pago.")
add_bullet("Si el cliente estaba suspendido, se reactiva automáticamente (pasa a Activo).")

add_heading("7.2. Historial de pagos", 2)
add_para("Debajo del formulario verá la tabla con todos los pagos registrados, ordenados del más "
         "reciente al más antiguo, mostrando: Fecha, Cliente, Período, Monto y Nota.")

# ============================================================
# 8. ESTADO DEL SERVICIO
# ============================================================
add_heading("8. Cómo funciona el estado del servicio", 1)
add_para("El sistema administra el corte y la reactivación del servicio de forma automática, "
         "siguiendo estas reglas:")
add_para("1. Al registrar un cliente:", bold=True)
add_bullet("Con 'primer mes gratis' → queda Activo por 30 días.")
add_bullet("Sin 'primer mes gratis' → queda Suspendido hasta su primer pago.")
add_para("2. Al registrar un pago:", bold=True)
add_bullet("El vencimiento se mueve 30 días hacia adelante.")
add_bullet("El cliente queda Activo.")
add_para("3. Cuando vence la fecha de pago:", bold=True)
add_bullet("Si un cliente pasa su fecha de vencimiento sin pagar, el sistema lo marca como Suspendido automáticamente.")
add_bullet("Ese cliente aparecerá en Pagos pendientes en el Dashboard.")
add_note("En resumen: un cliente al día está Activo; un cliente con pago vencido se suspende solo. "
         "Registrar el pago lo reactiva.")

# ============================================================
# 9. CERRAR SESIÓN
# ============================================================
add_heading("9. Cerrar sesión", 1)
add_para("Pulse el botón de cerrar sesión en la parte superior del panel. Volverá a la pantalla de "
         "inicio de sesión y su sesión quedará cerrada de forma segura.")
add_note("Cierre siempre la sesión si usa una computadora compartida.")

# ============================================================
# 10. PREGUNTAS FRECUENTES
# ============================================================
add_heading("10. Preguntas frecuentes", 1)

faqs = [
    ("¿Por qué un cliente aparece como Suspendido si pagó?",
     "Verifique que el pago se haya registrado en la sección Pagos con la fecha correcta. Al "
     "registrarlo, el cliente se reactiva automáticamente. También puede activarlo manualmente desde "
     "Editar cliente."),
    ("¿Puedo cambiar la tarifa de un cliente?",
     "Sí. Use el botón Editar en la tabla de clientes y modifique el campo de tarifa mensual."),
    ("El total cobrado del mes no coincide con lo que esperaba.",
     "El indicador 'Cobrado este mes' solo suma los pagos cuya fecha de pago cae dentro del mes "
     "actual. Revise las fechas en el historial de pagos."),
    ("Me cerró la sesión solo.",
     "Por seguridad, la sesión expira a las 12 horas. Solo debe volver a iniciar sesión."),
    ("¿Cuántos usuarios pueden usar el sistema?",
     "Actualmente el sistema funciona con un único usuario administrador con acceso total a todas las funciones."),
    ("Olvidé la contraseña / quiero cambiarla.",
     "Las credenciales las gestiona el responsable técnico en la configuración del sistema. "
     "Contáctelo para restablecerlas o cambiarlas."),
]
for q, a in faqs:
    add_para(q, bold=True, color=AZUL_CLARO)
    add_para(a)
    doc.add_paragraph()

# Pie final
doc.add_paragraph()
end = doc.add_paragraph()
end.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = end.add_run("Manual de usuario — Plataforma Nexus ISP")
r.italic = True
r.font.size = Pt(9)
r.font.color.rgb = GRIS

doc.save("Manual_Usuario_Nexus_ISP.docx")
print("Documento generado: Manual_Usuario_Nexus_ISP.docx")
